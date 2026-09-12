"""Export Project SIGNAL simulation and search results in portable formats."""

from __future__ import annotations

import csv
from datetime import datetime, timezone
import json
from pathlib import Path
from typing import Any, Iterable

import numpy as np

from signal_lab.experiment.clustering import ClusterObservation
from signal_lab.physics.engine import SimulationEngine


SUMMARY_FIELDS = (
    "project",
    "exported_at",
    "step",
    "seed",
    "particle_count",
    "world_width",
    "world_height",
    "species_count",
    "interaction_radius",
    "core_radius_ratio",
    "core_repulsion",
    "force_gain",
    "damping",
    "dt",
    "average_speed",
    "speed_variance",
    "cluster_count",
    "largest_cluster",
    "spatial_entropy",
    "directional_coherence",
    "species_0_fraction",
    "species_1_fraction",
    "species_2_fraction",
)


def _json_value(value: Any) -> Any:
    """Convert NumPy/scalar values into JSON and workbook-safe values."""
    if isinstance(value, np.ndarray):
        return [_json_value(item) for item in value.tolist()]
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    return value


def _current_metrics(engine: SimulationEngine, clusters: Iterable[ClusterObservation]) -> dict[str, float]:
    speeds = np.linalg.norm(engine.state.velocities, axis=1)
    norms = np.linalg.norm(engine.state.velocities, axis=1)
    grid = np.zeros((20, 20), dtype=np.int32)
    coords = (
        engine.state.positions
        / np.array([engine.config.width, engine.config.height])
        * 20
    ).astype(int) % 20
    np.add.at(grid, (coords[:, 1], coords[:, 0]), 1)
    probabilities = grid.ravel() / max(len(coords), 1)
    probabilities = probabilities[probabilities > 0]
    entropy = float(-np.sum(probabilities * np.log(probabilities)))
    cluster_list = list(clusters)
    proportions = np.bincount(engine.state.species, minlength=engine.genome.species_count)
    proportions = proportions / max(len(engine.state.species), 1)
    return {
        "average_speed": float(np.mean(speeds)),
        "speed_variance": float(np.var(speeds)),
        "cluster_count": float(len(cluster_list)),
        "largest_cluster": float(max((cluster.particle_count for cluster in cluster_list), default=0)),
        "spatial_entropy": entropy,
        "directional_coherence": float(np.linalg.norm(np.sum(engine.state.velocities, axis=0)) / max(np.sum(norms), 1e-12)),
        "species_0_fraction": float(proportions[0]) if len(proportions) > 0 else 0.0,
        "species_1_fraction": float(proportions[1]) if len(proportions) > 1 else 0.0,
        "species_2_fraction": float(proportions[2]) if len(proportions) > 2 else 0.0,
    }


def _cluster_rows(clusters: Iterable[ClusterObservation]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for cluster in clusters:
        row = cluster.to_dict()
        distribution = row.pop("species_distribution", [])
        row.pop("particle_ids", None)
        row.pop("history", None)
        for index, value in enumerate(distribution):
            row[f"species_{index}_fraction"] = value
        rows.append(_json_value(row))
    return rows


def _cluster_history_rows(clusters: Iterable[ClusterObservation]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for cluster in clusters:
        for point in cluster.history:
            row = {"cluster_id": cluster.cluster_id, **point}
            distribution = row.pop("species_distribution", [])
            for index, value in enumerate(distribution):
                row[f"species_{index}_fraction"] = value
            rows.append(_json_value(row))
    return rows


def _particle_rows(engine: SimulationEngine) -> list[dict[str, Any]]:
    speeds = np.linalg.norm(engine.state.velocities, axis=1)
    return [
        {
            "particle_id": int(identifier),
            "species": int(species),
            "x": float(position[0]),
            "y": float(position[1]),
            "vx": float(velocity[0]),
            "vy": float(velocity[1]),
            "speed": float(speed),
        }
        for identifier, species, position, velocity, speed in zip(
            engine.state.ids,
            engine.state.species,
            engine.state.positions,
            engine.state.velocities,
            speeds,
        )
    ]


def _time_series_rows(history: dict[str, Iterable[float]]) -> list[dict[str, Any]]:
    names = tuple(history)
    series = [list(values) for values in history.values()]
    return [
        {"sample": index + 1, **{name: float(values[index]) for name, values in zip(names, series) if index < len(values)}}
        for index in range(max((len(values) for values in series), default=0))
    ]


def _search_rows(search_results_path: Path | str | None = None) -> list[dict[str, Any]]:
    path = Path(search_results_path) if search_results_path is not None else Path("results/experiment_001_baseline/experiment_001_baseline.csv")
    if search_results_path is None and not path.exists():
        path = Path("results/search_results.csv")
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return [{key: value for key, value in row.items()} for row in csv.DictReader(handle)]


def collect_export_data(
    engine: SimulationEngine,
    clusters: Iterable[ClusterObservation],
    history: dict[str, Iterable[float]],
    events: Iterable[str],
    search_results_path: Path | str | None = None,
) -> dict[str, Any]:
    """Take one consistent, serializable snapshot for all export formats."""
    cluster_list = list(clusters)
    metrics = _current_metrics(engine, cluster_list)
    summary: dict[str, Any] = {
        "project": "Project SIGNAL - Experiment #1",
        "exported_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "step": int(engine.step_count),
        "seed": int(engine.config.seed),
        "particle_count": int(engine.state.count),
        "world_width": float(engine.config.width),
        "world_height": float(engine.config.height),
        "species_count": int(engine.genome.species_count),
        "interaction_radius": float(engine.genome.interaction_radius),
        "core_radius_ratio": float(engine.genome.core_radius_ratio),
        "core_repulsion": float(engine.genome.core_repulsion),
        "force_gain": float(engine.genome.force_gain),
        "damping": float(engine.genome.damping),
        "dt": float(engine.genome.dt),
        **metrics,
    }
    genome_rows = [
        {"parameter": "interaction_radius", "value": float(engine.genome.interaction_radius)},
        {"parameter": "core_radius_ratio", "value": float(engine.genome.core_radius_ratio)},
        {"parameter": "core_repulsion", "value": float(engine.genome.core_repulsion)},
        {"parameter": "force_gain", "value": float(engine.genome.force_gain)},
        {"parameter": "damping", "value": float(engine.genome.damping)},
        {"parameter": "dt", "value": float(engine.genome.dt)},
    ]
    matrix_rows = [
        {"from_species": row, **{f"to_species_{column}": float(engine.genome.interaction_matrix[row, column]) for column in range(engine.genome.species_count)}}
        for row in range(engine.genome.species_count)
    ]
    return {
        "summary": _json_value(summary),
        "genome": genome_rows,
        "interaction_matrix": matrix_rows,
        "clusters": _cluster_rows(cluster_list),
        "cluster_history": _cluster_history_rows(cluster_list),
        "time_series": _time_series_rows(history),
        "events": [{"index": index + 1, "event": event} for index, event in enumerate(events)],
        "particles": _particle_rows(engine),
        "search_results": _search_rows(search_results_path),
    }


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: Iterable[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(fieldnames or (rows[0].keys() if rows else []))
    if rows:
        for row in rows:
            for key in row:
                if key not in fields:
                    fields.append(key)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        if fields:
            writer.writeheader()
            writer.writerows({key: row.get(key, "") for key in fields} for row in rows)


def export_csv_bundle(data: dict[str, Any], output_dir: Path | str) -> list[Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    summary_path = output / "project_signal_summary.csv"
    _write_csv(summary_path, [{"metric": key, "value": value} for key, value in data["summary"].items()], ("metric", "value"))
    paths.append(summary_path)
    for key in ("genome", "interaction_matrix", "clusters", "cluster_history", "time_series", "events", "particles", "search_results"):
        path = output / f"project_signal_{key}.csv"
        _write_csv(path, data[key])
        paths.append(path)
    return paths


def export_json(data: dict[str, Any], path: Path | str) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(_json_value(data), indent=2), encoding="utf-8")
    return output


def export_xlsx(data: dict[str, Any], path: Path | str) -> Path:
    """Write an analysis-ready workbook with summary and supporting tables."""
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:  # pragma: no cover - depends on local install
        raise RuntimeError("Excel export requires openpyxl. Run: py -3 -m pip install -r requirements.txt") from error

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="163247")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(size=14, bold=True, color="163247")
    sheets = (
        ("Summary", [{"metric": key, "value": value} for key, value in data["summary"].items()]),
        ("Genome", data["genome"]),
        ("Matrix", data["interaction_matrix"]),
        ("Clusters", data["clusters"]),
        ("Cluster History", data["cluster_history"]),
        ("Time Series", data["time_series"]),
        ("Events", data["events"]),
        ("Particles", data["particles"]),
        ("Search Results", data["search_results"]),
    )
    for sheet_name, rows in sheets:
        sheet = workbook.create_sheet(sheet_name)
        if not rows:
            sheet.append(["No data"])
            sheet["A1"].font = header_font
            sheet["A1"].fill = header_fill
            continue
        fields = list(rows[0].keys())
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(field, "") for field in fields])
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(42, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        if sheet.max_row >= 2 and sheet_name != "Summary":
            table_name = "Table" + "".join(character for character in sheet_name if character.isalnum())
            table = Table(displayName=table_name[:250], ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
    summary_sheet = workbook["Summary"]
    summary_sheet["A1"].font = title_font
    summary_sheet.insert_rows(1)
    summary_sheet["A1"] = "Project SIGNAL - Experiment #1 export"
    summary_sheet["A1"].font = title_font
    summary_sheet.merge_cells("A1:B1")
    summary_sheet.freeze_panes = "A3"
    if len(data["time_series"]) >= 2:
        chart_sheet = workbook["Time Series"]
        chart = LineChart()
        chart.title = "Live metrics"
        chart.y_axis.title = "Value"
        chart.x_axis.title = "Sample"
        chart.add_data(Reference(chart_sheet, min_col=2, max_col=chart_sheet.max_column, min_row=1, max_row=chart_sheet.max_row), titles_from_data=True)
        chart.set_categories(Reference(chart_sheet, min_col=1, min_row=2, max_row=chart_sheet.max_row))
        chart.height = 8
        chart.width = 16
        chart_sheet.add_chart(chart, "J2")
    workbook.save(output)
    return output


def export_pdf(data: dict[str, Any], path: Path | str) -> Path:
    """Create a concise multi-page PDF report for sharing and review."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as error:  # pragma: no cover - depends on local install
        raise RuntimeError("PDF export requires reportlab. Run: py -3 -m pip install -r requirements.txt") from error

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SignalTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, textColor=colors.HexColor("#163247"), spaceAfter=10))
    styles.add(ParagraphStyle(name="SignalHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#163247"), spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle(name="SignalBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11))
    document = SimpleDocTemplate(str(output), pagesize=landscape(letter), rightMargin=0.45 * inch, leftMargin=0.45 * inch, topMargin=0.4 * inch, bottomMargin=0.4 * inch)
    story: list[Any] = [Paragraph("Project SIGNAL - Experiment #1", styles["SignalTitle"]), Paragraph("Simulation and structure export", styles["SignalBody"]), Spacer(1, 8)]

    summary_rows = [["Metric", "Value"]] + [[key.replace("_", " ").title(), str(value)] for key, value in data["summary"].items()]
    summary_table = Table(summary_rows, colWidths=[2.2 * inch, 2.0 * inch], repeatRows=1)
    summary_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163247")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C8D1")), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4F7")]), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story.extend([Paragraph("Run summary", styles["SignalHeading"]), summary_table])

    clusters = data["clusters"]
    if clusters:
        story.append(Paragraph("Detected structures", styles["SignalHeading"]))
        fields = ["cluster_id", "classification", "particle_count", "age_steps", "structure_score", "cohesion_score", "identity_score", "shape_score", "dynamic_score", "lifetime_score"]
        rows = [[field.replace("_", " ").title() for field in fields]] + [[str(row.get(field, "")) for field in fields] for row in clusters[:25]]
        table = Table(rows, repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163247")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C8D1")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4F7")])]))
        story.append(table)
    story.append(PageBreak())

    story.append(Paragraph("Search results", styles["SignalHeading"]))
    search_rows = data["search_results"]
    if search_rows:
        fields = list(search_rows[0].keys())
        rows = [[field.replace("_", " ").title() for field in fields]] + [[str(row.get(field, ""))[:52] for field in fields] for row in search_rows[:30]]
        table = Table(rows, repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163247")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C8D1")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4F7")])]))
        story.append(table)
    else:
        story.append(Paragraph("No search results CSV was available at export time.", styles["SignalBody"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph("Event log", styles["SignalHeading"]))
    event_rows = [["Index", "Event"]] + [[str(row["index"]), str(row["event"])] for row in data["events"][-35:]]
    table = Table(event_rows, colWidths=[0.6 * inch, 9.5 * inch], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163247")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C8D1")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4F7")]), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(table)
    document.build(story)
    return output


def export_all(data: dict[str, Any], output_dir: Path | str) -> dict[str, Path | list[Path] | str]:
    """Write the complete export bundle and return its paths."""
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    exported: dict[str, Path | list[Path]] = {
        "pdf": export_pdf(data, output / "project_signal_report.pdf"),
        "json": export_json(data, output / "project_signal_export.json"),
        "csv": export_csv_bundle(data, output),
    }
    try:
        exported["xlsx"] = export_xlsx(data, output / "project_signal_export.xlsx")
    except RuntimeError as error:
        exported["xlsx_error"] = str(error)
    return exported
